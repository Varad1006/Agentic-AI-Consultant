import io
from dataclasses import replace

import docx
import pytest

from app import documents
from app.documents import DocumentError, clean_text, extract_text, prepare_text


def test_text_files_are_decoded():
    assert extract_text("notes.TXT", "Hello\r\nworld".encode("utf-8-sig")) == "Hello\r\nworld"


def test_docx_paragraphs_and_tables_are_extracted():
    document = docx.Document()
    document.add_paragraph("Dispatch is planned in Excel.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Process"
    table.rows[0].cells[1].text = "Hours"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text("sop.docx", buffer.getvalue())
    assert "Dispatch is planned in Excel." in text
    assert "Process | Hours" in text


def test_unsupported_extension_is_rejected():
    with pytest.raises(DocumentError, match="Unsupported file type"):
        extract_text("image.png", b"...")


def test_invalid_utf8_is_rejected():
    with pytest.raises(DocumentError, match="UTF-8"):
        extract_text("notes.txt", b"\xff\xfe\xfa")


def test_clean_text_normalises_whitespace_and_control_characters():
    assert clean_text("a\x00b   c\n\n\n\nd \t e") == "ab c\n\nd e"


def test_pdf_word_breaks_are_rejoined():
    broken = "\n \n".join(f"word{i}" for i in range(40))
    fixed = documents._fix_pdf_word_breaks(broken)
    assert fixed.startswith("word0 word1 word2")


def test_prepare_text_rejects_short_documents():
    with pytest.raises(DocumentError, match="too short"):
        prepare_text("Too short.")


def test_prepare_text_truncates_long_documents(monkeypatch):
    monkeypatch.setattr(documents, "settings", replace(documents.settings, max_document_chars=100))
    text, truncated = prepare_text("word " * 100)
    assert truncated and len(text) == 100
