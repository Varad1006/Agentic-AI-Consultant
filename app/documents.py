"""Text extraction and clean-up for uploaded documents."""

import io
import re
from pathlib import Path

import docx
from pypdf import PdfReader

from app.config import settings

TEXT_EXTENSIONS = {".txt", ".md", ".csv"}
SUPPORTED_EXTENSIONS = {".pdf", ".docx"} | TEXT_EXTENSIONS
MIN_DOCUMENT_CHARS = 50

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f​﻿]")
# Some PDFs extract as "word\n \nword"; rejoin those words onto one line.
_PDF_WORD_BREAK = re.compile(r"(?<=\S)[ \t]*\n[ \t]*\n?[ \t]*(?=\S)")


class DocumentError(ValueError):
    """The document can't be used; the message is safe to show to the user."""


def extract_text(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentError(f"Unsupported file type '{extension or filename}'. Use one of: {supported}.")

    try:
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            pages = [_fix_pdf_word_breaks(page.extract_text() or "") for page in reader.pages]
            return "\n\n".join(pages)
        if extension == ".docx":
            document = docx.Document(io.BytesIO(data))
            paragraphs = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(paragraphs)
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentError("The file isn't valid UTF-8 text.") from exc
    except Exception as exc:
        raise DocumentError(f"Could not read the {extension[1:].upper()} file: {exc}") from exc


def prepare_text(text: str) -> tuple[str, bool]:
    """Clean the text and enforce size limits. Returns (text, was_truncated)."""
    text = clean_text(text)
    if len(text) < MIN_DOCUMENT_CHARS:
        raise DocumentError(
            "The document is too short to analyse. Include a description of the business and its processes."
        )
    if len(text) > settings.max_document_chars:
        return text[: settings.max_document_chars], True
    return text, False


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = _CONTROL_CHARS.sub("", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _fix_pdf_word_breaks(text: str) -> str:
    # Only rewrite pages where most line breaks are single-word lines (the extraction artifact);
    # normal pages keep their line structure.
    lines = [line for line in text.split("\n") if line.strip()]
    if len(lines) < 20:
        return text
    single_word_lines = sum(1 for line in lines if len(line.split()) <= 1)
    if single_word_lines / len(lines) < 0.6:
        return text
    return _PDF_WORD_BREAK.sub(" ", text)
